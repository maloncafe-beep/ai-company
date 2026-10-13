import glob
import os
import sys
import openpyxl

SRC_DIR = sys.argv[1] if len(sys.argv) > 1 else "kintai"
OUT = sys.argv[2] if len(sys.argv) > 2 else "集計まとめ.xlsx"
SHEET = "集計"

out_wb = openpyxl.Workbook()
out = out_wb.active
out.title = "集計まとめ"
log = out_wb.create_sheet("取込ログ")
log.append(["ファイル", "結果", "メモ"])

header_written = False
ok = ng = 0

for path in sorted(glob.glob(os.path.join(SRC_DIR, "*.xlsx"))):
    name = os.path.basename(path)
    if name.startswith("~$"):
        log.append([name, "対象外", "Excelの一時ファイル"])
        continue
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        target = next((s for s in wb.sheetnames if s.strip() == SHEET), None)
        if target is None:
            log.append([name, "失敗", f"「{SHEET}」シートが見つからない（あるシート: {', '.join(wb.sheetnames)}）"])
            ng += 1
            continue
        rows = list(wb[target].iter_rows(values_only=True))
        if len(rows) < 2:
            log.append([name, "失敗", "データ行がない"])
            ng += 1
            continue
        if not header_written:
            out.append(["元ファイル"] + list(rows[0]))
            header_written = True
        for r in rows[1:]:
            out.append([name] + list(r))
        log.append([name, "成功", f"{len(rows) - 1}行を取り込み" + ("（シート名の末尾の空白を補正）" if target != SHEET else "")])
        ok += 1
    except Exception as e:
        log.append([name, "失敗", f"{type(e).__name__}: {e}"])
        ng += 1

out_wb.save(OUT)
print(f"成功 {ok} / 失敗 {ng} → {OUT}")
