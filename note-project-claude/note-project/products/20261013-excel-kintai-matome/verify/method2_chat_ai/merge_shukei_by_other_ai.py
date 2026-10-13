# -*- coding: utf-8 -*-
"""
フォルダ内の複数のExcelファイル(.xlsx)から「集計」シートだけを取り出し、
1つのブックに縦に並べてまとめるスクリプトです。
使うライブラリは openpyxl のみです。
"""

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

DEFAULT_FOLDER = "kintai"
TARGET_SHEET = "集計"
OUTPUT_FILE = "まとめ.xlsx"
SOURCE_COLUMN_NAME = "元ファイル"
LOG_SHEET_NAME = "取込ログ"


def clean_value(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def find_target_sheet(workbook):
    for name in workbook.sheetnames:
        if name.strip() == TARGET_SHEET:
            return workbook[name]
    return None


def read_sheet_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = find_target_sheet(wb)
        if ws is None:
            raise ValueError(f"「{TARGET_SHEET}」シートがありません")

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError(f"「{TARGET_SHEET}」シートが空です")

        header = list(rows[0])
        while header and header[-1] is None:
            header.pop()

        data = []
        for row in rows[1:]:
            if all(cell is None or cell == "" for cell in row):
                continue
            data.append([clean_value(c) for c in row])
        return [clean_value(h) for h in header], data
    finally:
        wb.close()


def main():
    folder = Path(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FOLDER)
    if not folder.is_dir():
        print(f"フォルダが見つかりません: {folder}")
        print("フォルダ名を確認してください。")
        return

    out_path = Path(OUTPUT_FILE)

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "集計まとめ"
    log_ws = out_wb.create_sheet(LOG_SHEET_NAME)
    log_ws.append(["ファイル名", "結果", "メモ"])

    header_written = False
    base_header = []
    ok_count = 0
    ng_count = 0
    skip_count = 0

    files = sorted(
        p for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() == ".xlsx"
    )

    for path in files:
        name = path.name

        if name.startswith("~$"):
            log_ws.append([name, "対象外", "Excelの一時ファイルのため"])
            skip_count += 1
            continue

        if path.resolve() == out_path.resolve():
            log_ws.append([name, "対象外", "出力ファイルのため"])
            skip_count += 1
            continue

        try:
            header, data = read_sheet_rows(path)

            if not header_written:
                base_header = header
                out_ws.append([SOURCE_COLUMN_NAME] + header)
                header_written = True

            memo = f"{len(data)}行を取り込み"
            if header != base_header:
                memo += "（見出しが他のファイルと異なります。列のずれに注意）"

            for row in data:
                out_ws.append([name] + row)

            log_ws.append([name, "成功", memo])
            ok_count += 1

        except Exception as e:
            log_ws.append([name, "失敗", f"{type(e).__name__}: {e}"])
            ng_count += 1

    if not header_written:
        out_ws.append([SOURCE_COLUMN_NAME])

    log_ws.column_dimensions["A"].width = 40
    log_ws.column_dimensions["B"].width = 10
    log_ws.column_dimensions["C"].width = 60

    try:
        out_wb.save(out_path)
    except PermissionError:
        print(f"保存できませんでした。「{out_path}」をExcelで開いている場合は閉じて、もう一度実行してください。")
        return

    print("処理が終わりました。")
    print(f"  成功: {ok_count} 件")
    print(f"  失敗: {ng_count} 件")
    print(f"  対象外: {skip_count} 件")
    print(f"  出力: {out_path.resolve()}")
    print(f"  詳細は「{LOG_SHEET_NAME}」シートを見てください。")


if __name__ == "__main__":
    main()
