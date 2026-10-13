# -*- coding: utf-8 -*-
"""kintai フォルダ内の勤務簿(.xlsx)から「集計」シートだけを集めて result.xlsx にまとめる。
使い方: python merge.py [入力フォルダ] [出力ファイル]   (省略時はこのスクリプトと同じ場所の kintai / result.xlsx)
元ファイルは読み取り専用で開くだけで、変更しない。
"""
import sys, glob, os
from openpyxl import load_workbook, Workbook

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "kintai")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(BASE, "result.xlsx")
TARGET = "集計"

out = Workbook()
ws = out.active
ws.title = TARGET
log = []          # (ファイル名, 結果, 理由)
header = None

for path in sorted(glob.glob(os.path.join(SRC, "*.xlsx"))):
    name = os.path.basename(path)
    if name.startswith("~$"):
        log.append((name, "スキップ", "Excelの一時ファイル"))
        continue
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            # 前後の空白(全角含む)を除いて一致するシートを探す
            hit = next((s for s in wb.sheetnames if s.strip() == TARGET), None)
            if hit is None:
                log.append((name, "失敗", f"「{TARGET}」シートが無い(シート: {', '.join(wb.sheetnames)})"))
                continue
            rows = [r for r in wb[hit].iter_rows(values_only=True)
                    if any(c is not None for c in r)]
        finally:
            wb.close()
        if not rows:
            log.append((name, "失敗", "シートが空"))
            continue
        head, body = rows[0], rows[1:]
        if header is None:
            header = head
            ws.append(["ファイル名"] + list(head))
        note = ""
        if tuple(head) != tuple(header):
            note = "見出しが他ファイルと異なる(そのまま取込)"
        for r in body:
            ws.append([name] + list(r))
        log.append((name, "成功", f"{len(body)}行取込" + (" / " + note if note else "")
                    + (f" / 元シート名「{hit}」" if hit != TARGET else "")))
    except Exception as e:
        log.append((name, "失敗", f"開けない: {type(e).__name__}: {e}"))

lg = out.create_sheet("取込ログ")
lg.append(["ファイル名", "結果", "理由"])
for row in log:
    lg.append(list(row))
for sh in (ws, lg):
    for col in sh.columns:
        sh.column_dimensions[col[0].column_letter].width = min(
            60, max(len(str(c.value or "")) * 1.6 for c in col) + 2)
out.save(OUT)

ok = sum(1 for l in log if l[1] == "成功")
ng = sum(1 for l in log if l[1] == "失敗")
sk = sum(1 for l in log if l[1] == "スキップ")
print(f"成功 {ok} / 失敗 {ng} / スキップ {sk} -> {OUT}")
for l in log:
    if l[1] != "成功":
        print(" ", *l)
