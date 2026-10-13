# -*- coding: utf-8 -*-
"""
売上一覧のExcelを「取引先ごと」に別々のExcelファイルへ分割するスクリプト。

使うライブラリ: openpyxl のみ（あとは Python に最初から入っているもの）
"""

import os
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

TORIHIKISAKI_COL = "取引先"
KINGAKU_COL = "金額"
DEFAULT_OUTPUT_DIR = "取引先別"
UNSET_NAME = "取引先未設定"
LOG_FILE_NAME = "分割ログ"
MAX_NAME_LEN = 100

INVALID_CHARS = r'\/:*?"<>|'
RESERVED_NAMES = {"CON", "PRN", "AUX", "NUL"} | \
    {f"COM{i}" for i in range(1, 10)} | {f"LPT{i}" for i in range(1, 10)}


def is_blank(value):
    if value is None:
        return True
    return isinstance(value, str) and value.strip() == ""


def parse_amount(value):
    if is_blank(value):
        return Decimal(0), "blank"
    if isinstance(value, bool):
        return Decimal(0), "invalid"
    if isinstance(value, (int, float)):
        try:
            d = Decimal(str(value))
        except InvalidOperation:
            return Decimal(0), "invalid"
        return (d, "ok") if d.is_finite() else (Decimal(0), "invalid")
    if isinstance(value, str):
        s = unicodedata.normalize("NFKC", value).strip()
        negative = False
        if s[:1] in ("△", "▲"):
            negative, s = True, s[1:]
        if s.startswith("(") and s.endswith(")"):
            negative, s = True, s[1:-1]
        s = re.sub(r"[,¥円\s]", "", s)
        try:
            d = Decimal(s)
        except InvalidOperation:
            return Decimal(0), "invalid"
        if not d.is_finite():
            return Decimal(0), "invalid"
        return (-d if negative else d), "ok"
    return Decimal(0), "invalid"


def to_excel_number(d):
    if d == d.to_integral_value():
        return int(d)
    return float(d)


def sanitize_filename(name):
    notes = []
    result = name

    replaced = [c for c in INVALID_CHARS if c in result]
    if replaced:
        for c in replaced:
            result = result.replace(c, "_")
        notes.append("ファイル名に使えない文字（" + " ".join(replaced) + "）を「_」に置換")

    if re.search(r"[\x00-\x1f]", result):
        result = re.sub(r"[\x00-\x1f]", "_", result)
        notes.append("改行・タブなどの制御文字を「_」に置換")

    stripped = result.rstrip(" .")
    if stripped != result:
        result = stripped
        notes.append("末尾の「.」や空白を除去")

    if result == "":
        result = "_"
        notes.append("ファイル名が空になるため「_」に変更")

    if result.split(".")[0].upper() in RESERVED_NAMES:
        result = "_" + result
        notes.append("Windowsの予約名のため先頭に「_」を追加")

    if len(result) > MAX_NAME_LEN:
        result = result[:MAX_NAME_LEN].rstrip(" .")
        notes.append(f"{MAX_NAME_LEN}文字を超えるため切り詰め")

    return result, notes


def unique_filename(base, used_keys):
    candidate = base
    n = 2
    while candidate.lower() in used_keys:
        candidate = f"{base}_{n}"
        n += 1
    used_keys.add(candidate.lower())
    return candidate, candidate != base


def find_column(header_values, title):
    for i, h in enumerate(header_values):
        if isinstance(h, str):
            h = unicodedata.normalize("NFKC", h).strip()
        if h == title:
            return i
    return None


def main():
    args = sys.argv[1:]
    if args:
        input_path = args[0].strip().strip('"')
    else:
        input_path = input("分割したいExcelファイル(.xlsx)のパスを入力してください: ").strip().strip('"')
    output_dir = args[1].strip().strip('"') if len(args) >= 2 else DEFAULT_OUTPUT_DIR

    if not os.path.isfile(input_path):
        print(f"[エラー] ファイルが見つかりません: {input_path}")
        return 1
    if not input_path.lower().endswith((".xlsx", ".xlsm")):
        print("[エラー] .xlsx ファイルを指定してください。")
        return 1

    print(f"読み込み中: {input_path}")
    wb = load_workbook(input_path, data_only=True)
    ws = wb.worksheets[0]
    print(f"対象シート: {ws.title}")

    header = [c.value for c in ws[1]]
    col_name = find_column(header, TORIHIKISAKI_COL)
    col_amt = find_column(header, KINGAKU_COL)
    if col_name is None or col_amt is None:
        print(f"[エラー] 見出しに「{TORIHIKISAKI_COL}」「{KINGAKU_COL}」が見つかりません。")
        print(f"        見つかった見出し: {[h for h in header if h is not None]}")
        return 1
    n_cols = len(header)

    groups = {}
    unset = {"rows": [], "sum": Decimal(0), "blank_amt": 0, "bad_amt": 0}
    skipped_blank_rows = 0
    src_row_count = 0
    src_total = Decimal(0)

    for row in ws.iter_rows(min_row=2, max_col=n_cols):
        values = [c.value for c in row]
        if all(is_blank(v) for v in values):
            skipped_blank_rows += 1
            continue

        src_row_count += 1
        amount, status = parse_amount(values[col_amt])
        src_total += amount

        cells = [(c.value, c.number_format) for c in row]

        raw_name = values[col_name]
        if is_blank(raw_name):
            target = unset
        else:
            name = str(raw_name).strip()
            if name not in groups:
                groups[name] = {"rows": [], "sum": Decimal(0), "blank_amt": 0, "bad_amt": 0}
            target = groups[name]

        target["rows"].append(cells)
        target["sum"] += amount
        if status == "blank":
            target["blank_amt"] += 1
        elif status == "invalid":
            target["bad_amt"] += 1

    if src_row_count == 0:
        print("[エラー] データ行がありません（2行目以降が空です）。")
        return 1

    os.makedirs(output_dir, exist_ok=True)

    used_keys = {LOG_FILE_NAME.lower()}
    plan = []

    if unset["rows"]:
        used_keys.add(UNSET_NAME.lower())
        plan.append(("(空欄)", UNSET_NAME, unset,
                     [f"取引先が空欄の行 {len(unset['rows'])} 件をまとめました"]))

    for name, info in groups.items():
        base, notes = sanitize_filename(name)
        fname, renamed = unique_filename(base, used_keys)
        if renamed:
            notes.append(f"同じファイル名の取引先があるため「{fname}」に変更")
        plan.append((name, fname, info, notes))

    log_rows = []
    file_paths = []

    for disp_name, fname, info, notes in plan:
        out_wb = Workbook()
        out_ws = out_wb.active
        out_ws.title = ws.title[:31]

        out_ws.append(header)
        for i in range(1, n_cols + 1):
            letter = get_column_letter(i)
            w = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
            if w:
                out_ws.column_dimensions[letter].width = w

        for r_idx, cells in enumerate(info["rows"], start=2):
            for c_idx, (val, fmt) in enumerate(cells, start=1):
                cell = out_ws.cell(row=r_idx, column=c_idx, value=val)
                if fmt and fmt != "General":
                    cell.number_format = fmt

        out_path = os.path.join(output_dir, fname + ".xlsx")
        try:
            out_wb.save(out_path)
        except PermissionError:
            print(f"[エラー] 保存できません: {out_path}")
            print("        同名のExcelファイルを開いたままではありませんか？ 閉じてからやり直してください。")
            return 1
        file_paths.append((out_path, info))

        if info["blank_amt"]:
            notes.append(f"金額が空欄の行 {info['blank_amt']} 件（0として集計）")
        if info["bad_amt"]:
            notes.append(f"金額が数値でない行 {info['bad_amt']} 件（0として集計）")
        log_rows.append((disp_name, fname + ".xlsx", len(info["rows"]),
                         to_excel_number(info["sum"]), "／".join(notes)))

        for n in notes:
            print(f"  [注意] {disp_name}: {n}")

    out_row_count = 0
    out_total = Decimal(0)
    for out_path, _info in file_paths:
        chk_ws = load_workbook(out_path, data_only=True).worksheets[0]
        for r in chk_ws.iter_rows(min_row=2, values_only=True):
            if all(is_blank(v) for v in r):
                continue
            out_row_count += 1
            amt, _ = parse_amount(r[col_amt])
            out_total += amt

    rows_ok = (src_row_count == out_row_count)
    total_ok = (src_total == out_total)
    verdict = "OK（一致）" if (rows_ok and total_ok) else "NG（不一致：要確認）"

    print("")
    print("========== 検算結果 ==========")
    print(f"元ファイル   : {src_row_count} 行 / 金額合計 {src_total:,}")
    print(f"分割後の合計 : {out_row_count} 行 / 金額合計 {out_total:,}")
    print(f"行数 : {'一致' if rows_ok else '不一致'}　金額 : {'一致' if total_ok else '不一致'}")
    print(f"判定 : {verdict}")
    if skipped_blank_rows:
        print(f"（完全に空の行 {skipped_blank_rows} 行は対象外として読み飛ばしました）")
    print(f"出力ファイル数: {len(file_paths)} 個 → フォルダ「{output_dir}」")

    log_wb = Workbook()
    log_ws = log_wb.active
    log_ws.title = "分割ログ"
    log_ws.append(["取引先", "ファイル名", "件数", "金額合計", "メモ"])
    for r in log_rows:
        log_ws.append(list(r))

    log_ws.append([])
    log_ws.append(["【分割後の合計】", f"{len(file_paths)} ファイル", out_row_count,
                   to_excel_number(out_total), ""])
    log_ws.append(["【元ファイル】", os.path.basename(input_path), src_row_count,
                   to_excel_number(src_total),
                   f"完全に空の行 {skipped_blank_rows} 行は対象外" if skipped_blank_rows else ""])
    log_ws.append(["【検算】", "", "一致" if rows_ok else "不一致",
                   "一致" if total_ok else "不一致", f"判定: {verdict}"])

    for letter, width in zip("ABCDE", (28, 34, 8, 16, 70)):
        log_ws.column_dimensions[letter].width = width
    for r in range(2, log_ws.max_row + 1):
        log_ws.cell(row=r, column=4).number_format = "#,##0.##"

    log_path = os.path.join(output_dir, LOG_FILE_NAME + ".xlsx")
    try:
        log_wb.save(log_path)
    except PermissionError:
        print(f"[エラー] ログを保存できません: {log_path}（開いたままのExcelを閉じてください）")
        return 1
    print(f"分割ログ: {log_path}")

    return 0 if (rows_ok and total_ok) else 1


if __name__ == "__main__":
    try:
        code = main()
    except Exception as e:
        print(f"[エラー] 予期しない問題が起きました: {e}")
        code = 1
    if sys.stdin and sys.stdin.isatty():
        input("\nEnterキーを押すと終了します。")
    sys.exit(code)
