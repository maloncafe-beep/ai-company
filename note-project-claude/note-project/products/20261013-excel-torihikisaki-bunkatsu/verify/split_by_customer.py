import os
import re
import sys
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "売上一覧_2026-09.xlsx"
OUT_DIR = sys.argv[2] if len(sys.argv) > 2 else "取引先別"
KEY_COL = "取引先"      # 分けたい列の見出し
AMOUNT_COL = "金額"     # 検算に使う列の見出し

os.makedirs(OUT_DIR, exist_ok=True)

ws = openpyxl.load_workbook(SRC, data_only=True).worksheets[0]
rows = list(ws.iter_rows(values_only=True))
header, data = list(rows[0]), [r for r in rows[1:] if any(v is not None for v in r)]
key_idx, amt_idx = header.index(KEY_COL), header.index(AMOUNT_COL)

# 取引先ごとにまとめる（前後の空白は取り除く。空欄は「取引先未設定」）
groups = {}
for r in data:
    key = str(r[key_idx]).strip() if r[key_idx] is not None else ""
    groups.setdefault(key or "取引先未設定", []).append(list(r))

log = openpyxl.Workbook()
lg = log.active
lg.title = "分割ログ"
lg.append(["取引先", "ファイル名", "件数", "金額合計", "メモ"])

used = set()
written = []
for key, items in sorted(groups.items()):
    memo = []
    base = re.sub(r'[\\/:*?"<>|]', "_", key)
    if base != key:
        memo.append("ファイル名に使えない文字を「_」に置換")
    fname, n = base, 2
    while fname.lower() in used:          # 置換後に名前が重なったら連番を付ける
        fname = f"{base}_{n}"
        n += 1
    if fname != base:
        memo.append("同じ名前のファイルがあるため連番を付けた")
    used.add(fname.lower())
    if key == "取引先未設定":
        memo.append("取引先が空欄の行。要確認")

    out = openpyxl.Workbook()
    o = out.active
    o.title = "売上"
    o.append(header)
    for r in items:
        o.append(r)
    out.save(os.path.join(OUT_DIR, fname + ".xlsx"))
    written.append(fname + ".xlsx")
    lg.append([key, fname + ".xlsx", len(items), sum(r[amt_idx] or 0 for r in items), " / ".join(memo)])

# 検算：書き出したファイルを読み直して、元の行数・金額と比べる
out_rows = out_amount = 0
for f in written:
    w = openpyxl.load_workbook(os.path.join(OUT_DIR, f)).worksheets[0]
    for r in w.iter_rows(min_row=2, values_only=True):
        out_rows += 1
        out_amount += r[amt_idx] or 0
src_rows = len(data)
src_amount = sum(r[amt_idx] or 0 for r in data)
ok = (src_rows == out_rows) and (src_amount == out_amount)

lg.append([])
lg.append(["検算", "", f"元 {src_rows} / 分割後 {out_rows}", f"元 {src_amount} / 分割後 {out_amount}", "一致" if ok else "不一致。要確認"])
log.save(os.path.join(OUT_DIR, "_分割ログ.xlsx"))

print(f"{len(groups)} 社に分割 / 元 {src_rows} 行 → 分割後 {out_rows} 行 / 検算 {'OK' if ok else 'NG（要確認）'}")
