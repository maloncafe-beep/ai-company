"""
クラウドワークス 案件検索 → Excel出力スクリプト（Selenium版 v3）
事前準備:
  pip install selenium openpyxl
"""

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoSuchElementException
import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import time
import re
import traceback
from urllib.parse import quote


def extract_deadline(deadline_text):
    """
    「あと3日（6月28日まで）」のようなテキストから期限日を抽出
    複数のパターンに対応
    """
    if not deadline_text:
        return ""

    # パターン1: 「X月X日」形式を抽出
    match = re.search(r'(\d{1,2})月(\d{1,2})日', deadline_text)
    if match:
        month = match.group(1)
        day = match.group(2)
        return f"{month}月{day}日"

    # パターン2: 「あとX日」形式のみの場合
    match = re.search(r'あと(\d+)日', deadline_text)
    if match:
        days = match.group(1)
        return f"あと{days}日"

    # 上記パターンに当てはまらない場合は元のテキストを返す
    return deadline_text.strip()


# =============================
# 検索条件（ここを変更してください）
# =============================
SEARCH_CONFIG = {
    "keywords": ["提案書作成", "営業資料作成", "パワーポイント作成"],  # 複数キーワード指定可
    "employment_type": "fixed_worked",  # time_worked=時間単価, fixed_worked=プロジェクト
    "min_hourly_wage": 0,               # プロジェクト案件なので0
    "inexperienced": 0,                 # 0=すべて
    "order": "created_at",              # created_at=新着順（精度重視）
    "hide_expired": 1,
    "max_pages": 3,                     # 取得ページ数（1ページ約30件）
}

import os
import subprocess
import json

# tools/job-scout/ から2階層上がってプロジェクトルート → logs/ へ
_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
LOGS_DIR = os.path.join(_ROOT, "logs", "Case_list_py")
os.makedirs(LOGS_DIR, exist_ok=True)


def load_keywords_from_sheet():
    """
    node server/scripts/read-keywords.mjs を呼び出してスプシのキーワード設定を取得する。
    失敗時は SEARCH_CONFIG のデフォルト値にフォールバックする。
    """
    mjs_path = os.path.join(_ROOT, "server", "scripts", "read-keywords.mjs")
    try:
        result = subprocess.run(
            ["node", mjs_path],
            capture_output=True, text=True, timeout=30,
            cwd=_ROOT
        )
        if result.returncode != 0:
            print(f"⚠️  read-keywords.mjs エラー: {result.stderr.strip()}")
            print("   → SEARCH_CONFIG のデフォルトキーワードで続行します。")
            return None
        # dotenvx が stdout に余分なテキストを出力するため、行頭が [ または { の行から抽出する
        lines = result.stdout.splitlines()
        json_lines = []
        in_json = False
        for line in lines:
            if not in_json and (line.startswith('[') or line.startswith('{')):
                in_json = True
            if in_json:
                json_lines.append(line)
        data = json.loads('\n'.join(json_lines))
        return data
    except Exception as e:
        print(f"⚠️  キーワード取得失敗 ({e}) → デフォルトで続行します。")
        return None


# スプシからキーワードを読み込む（失敗時はデフォルト）
_sheet_keywords = load_keywords_from_sheet()

# スプシ → CSV を最新化（スプシで削除された案件をCSVから除去）
_dl_path = os.path.join(_ROOT, "server", "scripts", "download-sheet.mjs")
_dl = subprocess.run(["node", _dl_path], capture_output=True, text=True, timeout=30, cwd=_ROOT)
if _dl.returncode == 0:
    print("✅ スプシ → CSV 同期完了（削除済み案件をCSVから除去）")
else:
    print(f"⚠️  スプシ同期失敗（既存CSVで続行）: {_dl.stderr.strip()}")

# スプシカテゴリ → CrowdWorks group URL のマッピング
CATEGORY_GROUP_MAP = {
    "ライティング系":  "writing_beginner",
    "資料系":          "business",
    "文書・業務系":    "business",
    "SNS・マーケ系":   "business",
    "動画・台本系":    "video_contents",
}

if _sheet_keywords:
    # D列が有効なキーワードを (group, keyword) のペアで管理
    _search_pairs = []
    for k in _sheet_keywords:
        group = CATEGORY_GROUP_MAP.get(k["category"])
        if group:
            _search_pairs.append({"group": group, "keyword": k["keyword"], "category": k.get("category", "")})
        else:
            # マッピングにないカテゴリは従来の全体検索にフォールバック
            _search_pairs.append({"group": None, "keyword": k["keyword"], "category": k.get("category", "")})

    _search_kws = [p["keyword"] for p in _search_pairs]
    _kw_excludes = list({e for k in _sheet_keywords for e in k.get("excludes", [])})

    SEARCH_CONFIG["keywords"] = _search_kws if _search_kws else SEARCH_CONFIG["keywords"]
    SEARCH_CONFIG["search_pairs"] = _search_pairs
    print(f"✅ スプシからキーワード読み込み: {len(_search_kws)} 件（カテゴリ絞り込みあり）")
    if _kw_excludes:
        print(f"   キーワード別除外条件: {_kw_excludes}")
else:
    _search_pairs = [
        {"group": None, "keyword": kw, "category": ""}
        for kw in SEARCH_CONFIG["keywords"]
    ]
    _kw_excludes = []

# OUTPUT_FILE_XLSX は main() 内でモード選択後に確定する（モジュールレベルでは仮置き）
OUTPUT_FILE_XLSX = os.path.join(LOGS_DIR, f"crowdworks_pending_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
OUTPUT_FILE_MD = os.path.join(LOGS_DIR, "crowdworks_cases_latest.md")


# =============================
# モード選択（インタラクティブ）
# =============================

def select_mode():
    """
    実行モードをインタラクティブに選択する。
    戻り値: 選択された search_pairs (list of dict) と ラベル文字列のタプル
    """
    # 全ペアを取得
    all_pairs = SEARCH_CONFIG.get("search_pairs") or _search_pairs

    # カテゴリ一覧を重複なく抽出（スプシ由来）
    categories = []
    if _sheet_keywords:
        seen = set()
        for k in _sheet_keywords:
            cat = k.get("category", "")
            if cat and cat not in seen:
                categories.append(cat)
                seen.add(cat)

    while True:
        print()
        print("=" * 40)
        print("=== CrowdWorks 案件検索 ===")
        print("実行モードを選んでください：")
        print()
        print(f"  1. 全キーワード（{len(all_pairs)}件・通常モード）")
        print("  2. カテゴリ指定")
        print("  3. キーワード指定")
        print()

        mode = input("番号を入力 > ").strip()

        if mode == "1":
            selected_pairs = all_pairs
            label = f"全キーワード（{len(all_pairs)}件）"

        elif mode == "2":
            if not categories:
                print("⚠️  カテゴリ情報が取得できません（スプシ接続を確認してください）。")
                continue
            print()
            print("カテゴリ一覧：")
            for i, cat in enumerate(categories, start=1):
                kw_count = sum(1 for k in _sheet_keywords if k.get("category") == cat)
                print(f"  {i}. {cat}（{kw_count}キーワード）")
            print()
            cat_input = input("番号を入力 > ").strip()
            try:
                cat_idx = int(cat_input) - 1
                if cat_idx < 0 or cat_idx >= len(categories):
                    print("⚠️  無効な番号です。")
                    continue
                selected_cat = categories[cat_idx]
                selected_pairs = [
                    p for p in all_pairs
                    if p.get("category") == selected_cat or any(
                        k.get("keyword") == p["keyword"] and k.get("category") == selected_cat
                        for k in (_sheet_keywords or [])
                    )
                ]
                if not selected_pairs:
                    print(f"⚠️  カテゴリ「{selected_cat}」に対応するキーワードが見つかりません。")
                    continue
                label = f"カテゴリ「{selected_cat}」（{len(selected_pairs)}キーワード）"
            except ValueError:
                print("⚠️  数字を入力してください。")
                continue

        elif mode == "3":
            print()
            kw_input = input("キーワードを入力（複数はスペース区切り） > ").strip()
            if not kw_input:
                print("⚠️  キーワードを入力してください。")
                continue
            keywords = kw_input.split()
            selected_pairs = [{"group": None, "keyword": kw, "category": ""} for kw in keywords]
            label = f"キーワード指定：{' / '.join(keywords)}"

        else:
            print("⚠️  1・2・3 のいずれかを入力してください。")
            continue

        # 確認プロンプト
        print()
        print(f"以下の条件で実行します：")
        print(f"  {label}")
        kws = [p["keyword"] for p in selected_pairs]
        if len(kws) <= 8:
            for kw in kws:
                print(f"  ・{kw}")
        else:
            for kw in kws[:5]:
                print(f"  ・{kw}")
            print(f"  ・... 他 {len(kws) - 5} 件")
        print()
        confirm = input("よろしいですか？ [y/n] > ").strip().lower()
        if confirm == "y":
            return selected_pairs, label
        else:
            print("最初の選択に戻ります。")
            continue


def build_url(keyword, page=1, group=None):
    p = SEARCH_CONFIG
    if group:
        base = f"https://crowdworks.jp/public/jobs/group/{group}"
    else:
        base = "https://crowdworks.jp/public/jobs/search"
    params = (
        f"?employment_type={p['employment_type']}"
        f"&keyword={quote(keyword)}"
        f"&order={p['order']}"
        f"&hide_expired={p['hide_expired']}"
        f"&page={page}"
    )
    if p.get("min_hourly_wage"):
        params += f"&min_hourly_wage={p['min_hourly_wage']}"
    if p.get("inexperienced"):
        params += f"&inexperienced={p['inexperienced']}"
    return base + params


def init_driver():
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--window-size=1920,1080")
    options.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )
    return webdriver.Chrome(options=options)


def parse_jobs(driver):
    jobs = []

    try:
        WebDriverWait(driver, 15).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "a[class*='_titleLinkPc']"))
        )
    except TimeoutException:
        print("  タイムアウト: 案件タイトルが見つかりませんでした")
        return jobs

    time.sleep(1)

    title_links = driver.find_elements(By.CSS_SELECTOR, "a[class*='_titleLinkPc']")
    print(f"  案件タイトル数: {len(title_links)} 件")

    for title_link in title_links:
        job = {}
        job["タイトル"] = title_link.text.strip()
        job["URL"] = title_link.get_attribute("href")

        if not job["タイトル"] or not job["URL"]:
            continue

        m = re.search(r"/jobs/(\d+)", job["URL"])
        job["案件ID"] = m.group(1) if m else ""

        try:
            card = driver.execute_script("""
                let el = arguments[0];
                for (let i = 0; i < 6; i++) {
                    el = el.parentElement;
                    if (!el) break;
                    if (el.tagName === 'LI' || el.tagName === 'ARTICLE') return el;
                }
                return arguments[0].parentElement.parentElement.parentElement;
            """, title_link)

            try:
                wage_el = card.find_element(By.CSS_SELECTOR, "[class*='_wage'], [class*='_price'], [class*='_payment'], [class*='_salary']")
                job["単価"] = wage_el.text.strip()
            except NoSuchElementException:
                job["単価"] = ""

            try:
                client_el = card.find_element(By.CSS_SELECTOR, "[class*='_userName'], [class*='_client'], [class*='_employer']")
                job["クライアント"] = client_el.text.strip()
            except NoSuchElementException:
                job["クライアント"] = ""

            try:
                # 期限情報を複数の方法で取得を試みる
                deadline = ""

                # 方法1: より詳細なセレクタで期限を探す
                try:
                    # クラウドワークスの期限表示用クラスをターゲット
                    deadline_el = card.find_element(By.CSS_SELECTOR,
                        "[class*='_period'], [class*='_deadline'], [class*='_limit'], "
                        "[class*='_dueDate'], [class*='_endDate'], [class*='Deadline']")
                    deadline = deadline_el.text.strip()
                except NoSuchElementException:
                    # 方法2: カード内の全テキストを検索して「日」で終わる部分を探す
                    try:
                        card_text = card.text
                        # 「X月X日」パターンを全て抽出
                        matches = re.findall(r'(\d{1,2})月(\d{1,2})日', card_text)
                        if matches:
                            # 最初にマッチした日付を使用
                            month, day = matches[0]
                            deadline = f"{month}月{day}日"
                    except:
                        pass

                # 抽出した期限情報を整形
                job["締切"] = extract_deadline(deadline) if deadline else ""
            except Exception as e:
                print(f"    警告: 締切取得エラー - {e}")
                job["締切"] = ""

            try:
                apply_el = card.find_element(By.CSS_SELECTOR, "[class*='_apply'], [class*='_proposal'], [class*='_applicant']")
                job["応募数"] = apply_el.text.strip()
            except NoSuchElementException:
                job["応募数"] = ""

            try:
                tag_els = card.find_elements(By.CSS_SELECTOR, "[class*='_tagsLink'], [class*='_tag'], [class*='_skill']")
                job["タグ"] = " / ".join(t.text.strip() for t in tag_els[:5] if t.text.strip())
            except Exception:
                job["タグ"] = ""

        except Exception:
            job["単価"] = ""
            job["クライアント"] = ""
            job["締切"] = ""
            job["応募数"] = ""
            job["タグ"] = ""

        jobs.append(job)

    return jobs


def scrape_all():
    print("Chromeを起動中...")
    driver = init_driver()
    all_jobs = []

    try:
        search_pairs = SEARCH_CONFIG.get("search_pairs") or [
            {"group": None, "keyword": kw}
            for kw in (SEARCH_CONFIG["keywords"] if isinstance(SEARCH_CONFIG["keywords"], list) else [SEARCH_CONFIG["keywords"]])
        ]

        for keyword_idx, pair in enumerate(search_pairs, start=1):
            keyword = pair["keyword"]
            group = pair.get("group")
            print(f"\n{'='*50}")
            print(f"キーワード {keyword_idx}/{len(search_pairs)}: 「{keyword}」（group={group}）を検索中...")
            print(f"{'='*50}")

            for page in range(1, SEARCH_CONFIG["max_pages"] + 1):
                url = build_url(keyword, page, group=group)
                print(f"\n取得中: ページ {page}")
                print(f"  URL: {url}")
                driver.get(url)

                jobs = parse_jobs(driver)

                if not jobs:
                    print(f"  ページ {page} で案件が見つかりませんでした。次のキーワードに進みます。")
                    break

                all_jobs.extend(jobs)
                print(f"  {len(jobs)} 件取得（{keyword}の累計 {len([j for j in all_jobs if '案件ID' in j])} 件）")

                if page < SEARCH_CONFIG["max_pages"]:
                    time.sleep(2)

            if keyword_idx < len(search_pairs):
                time.sleep(2)

    finally:
        driver.quit()
        print("\nブラウザを終了しました。")

    return all_jobs


def save_to_markdown(jobs, filepath):
    """Markdown 形式で見やすく出力"""
    md_lines = [
        "# CrowdWorks 案件リスト（見やすい版）\n",
        f"**取得日時**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n",
        f"**取得件数**: {len(jobs)} 件\n",
        "---\n",
    ]

    md_lines.append("## 案件一覧\n")
    md_lines.append("| # | タイトル | 単価 | クライアント | 締切 | 応募数 | URL |\n")
    md_lines.append("|---|------|------|----------|------|--------|------|\n")

    for idx, job in enumerate(jobs[:20], start=1):  # 最初の20件
        title = job.get("タイトル", "").replace("|", "\\|")[:50]
        wage = job.get("単価", "")
        client = job.get("クライアント", "")[:20]
        deadline = job.get("締切", "")
        applies = job.get("応募数", "")
        url = job.get("URL", "")
        md_lines.append(f"| {idx} | {title} | {wage} | {client} | {deadline} | {applies} | [Link]({url}) |\n")

    with open(filepath, "w", encoding="utf-8") as f:
        f.writelines(md_lines)


def save_to_csv(jobs, csv_path):
    """
    logs/案件検索結果.csv の形式（11列）に変換して追記する。
    既存IDと重複する案件はスキップ。
    """
    # 既存IDを読み込んで重複チェック用セットを作る
    existing_ids = set()
    if os.path.exists(csv_path):
        with open(csv_path, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            next(reader, None)  # ヘッダーをスキップ
            for row in reader:
                if row:
                    existing_ids.add(row[0].strip())

    today = datetime.now().strftime("%Y/%m/%d")
    new_rows = []

    # スクリプト固定の除外 + スプシ由来の除外を統合
    TITLE_EXCLUDE = ['モデル', '面談', 'チャット', 'ライブ配信'] + _kw_excludes

    for job in jobs:
        job_id = job.get("案件ID", "").strip()
        if not job_id or job_id in existing_ids:
            continue

        title_raw = job.get("タイトル", "")
        if any(w in title_raw for w in TITLE_EXCLUDE):
            continue

        def clean(val):
            """改行・カンマ（桁区切り）・余分な空白を除去。クライアント名に混入する「掲載日：」以降も除去"""
            v = (val or "").replace("\n", " ").replace("\r", "").strip()
            v = re.sub(r"掲載日[：:].+", "", v).strip()
            v = re.sub(r"(気になる|この仕事に似た仕事を依頼する).*", "", v).strip()
            return re.sub(r"(?<=\d),(?=\d)", "", v)

        # 報酬からカンマ・改行を除去
        wage = clean(job.get("単価", "")) or "要確認"

        # 締切をYYYY/MM/DD形式に変換（「8月24日」→ 当年で補完）
        deadline_raw = job.get("締切", "") or ""
        deadline = ""
        m = re.search(r"(\d{1,2})月(\d{1,2})日", deadline_raw)
        if m:
            year = datetime.now().year
            deadline = f"{year}/{int(m.group(1)):02d}/{int(m.group(2)):02d}"
        if not deadline:
            deadline = "2026/09/30"

        applies = clean(job.get("応募数", ""))
        memo = f"応募{applies}" if applies else ""

        row = [
            job_id,                                      # A: 管理ID
            "ＣＷ",                                       # B: 媒体
            clean(job.get("タイトル", "")),               # C: 案件タイトル
            wage,                                         # D: 報酬
            clean(job.get("クライアント", "")) or "不明", # E: クライアント名
            deadline,                                     # F: 納期
            job.get("URL", ""),                           # G: URL
            memo,                                         # H: メモ
            "",                                           # I: 検討（社長記入欄）
            "",                                           # J: 仕事の詳細
            "",                                           # K: Masterメモ
        ]
        new_rows.append(row)
        existing_ids.add(job_id)

    if new_rows:
        with open(csv_path, "a", encoding="utf-8-sig", newline="\n") as f:
            writer = csv.writer(f, lineterminator="\n")
            writer.writerows(new_rows)
        print(f"✅ CSV追記完了: {len(new_rows)} 件 → {csv_path}")
    else:
        print("ℹ️  新規案件なし（全件重複またはID未取得）")

    return len(new_rows)


def save_to_excel(jobs):
    print(f"\nExcel出力開始... ({len(jobs)} 件)")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "案件一覧"

    columns = ["案件ID", "タイトル", "単価", "クライアント", "締切", "応募数", "タグ", "URL"]
    col_widths = [10, 45, 20, 25, 15, 10, 35, 60]

    header_fill = PatternFill("solid", start_color="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)

    print("  ヘッダー行を作成中...")
    for col_idx, (col_name, width) in enumerate(zip(columns, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    alt_fill = PatternFill("solid", start_color="EBF3FB")
    url_font = Font(color="1F4E79", underline="single", name="Arial", size=9)
    default_font = Font(name="Arial", size=9)

    print("  データ行を書き込み中...")
    for row_idx, job in enumerate(jobs, start=2):
        if row_idx % 20 == 0:
            print(f"    {row_idx - 1} 行目まで完了...")
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, col_name in enumerate(columns, start=1):
            val = job.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="center", wrap_text=(col_name == "タイトル"))
            if fill:
                cell.fill = fill
            if col_name == "URL" and val:
                cell.hyperlink = val
                cell.font = url_font
            else:
                cell.font = default_font

    # 検索条件シート
    print("  検索条件シートを作成中...")
    ws2 = wb.create_sheet("検索条件")
    ws2["A1"] = "検索条件"
    ws2["A1"].font = Font(bold=True, size=12, name="Arial")
    p = SEARCH_CONFIG
    keywords = p["keywords"] if isinstance(p["keywords"], list) else [p["keywords"]]
    conditions = [
        ("キーワード", " / ".join(keywords)),
        ("案件種別", "時間単価" if p["employment_type"] == "time_worked" else "プロジェクト"),
        ("時間単価（下限）", f"{p['min_hourly_wage']}円" if p.get("min_hourly_wage") else "指定なし"),
        ("未経験可", "はい" if p.get("inexperienced") else "すべて"),
        ("取得ページ数", str(p["max_pages"])),
        ("取得日時", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("取得件数", str(len(jobs))),
    ]
    for r, (k, v) in enumerate(conditions, start=2):
        ws2.cell(row=r, column=1, value=k).font = Font(bold=True, name="Arial", size=10)
        ws2.cell(row=r, column=2, value=v).font = Font(name="Arial", size=10)
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 30

    print(f"  ファイル保存中: {OUTPUT_FILE_XLSX}")
    wb.save(OUTPUT_FILE_XLSX)
    print(f"\n✅ Excel 保存完了: {OUTPUT_FILE_XLSX}（{len(jobs)} 件）")

    # Markdown 形式でも出力
    print(f"📝 Markdown 形式でも出力中...")
    save_to_markdown(jobs, OUTPUT_FILE_MD)
    print(f"✅ Markdown 保存完了: {OUTPUT_FILE_MD}")


def main():
    global OUTPUT_FILE_XLSX

    try:
        # ========== モード選択 ==========
        selected_pairs, mode_label = select_mode()

        # 選択結果を SEARCH_CONFIG に反映
        SEARCH_CONFIG["search_pairs"] = selected_pairs
        SEARCH_CONFIG["keywords"] = [p["keyword"] for p in selected_pairs]

        # 出力ファイル名をモード選択後のキーワードで確定
        keyword_str = "_".join(SEARCH_CONFIG["keywords"][:3])
        OUTPUT_FILE_XLSX = os.path.join(
            LOGS_DIR,
            f"crowdworks_{keyword_str}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )

        # ========== 実行サマリ表示 ==========
        print()
        print("=" * 50)
        print("クラウドワークス案件検索（Selenium版 v3）")
        print(f"モード       : {mode_label}")
        keywords = SEARCH_CONFIG["keywords"]
        if len(keywords) <= 5:
            print(f"キーワード   : {' / '.join(keywords)}")
        else:
            print(f"キーワード   : {' / '.join(keywords[:5])} ... 他{len(keywords)-5}件")
        print(f"時間単価     : {SEARCH_CONFIG.get('min_hourly_wage', '指定なし')}円以上")
        print(f"未経験可     : {'はい' if SEARCH_CONFIG.get('inexperienced') else 'すべて'}")
        print("=" * 50)

        jobs = scrape_all()

        if not jobs:
            print("\n❌ 案件が取得できませんでした。")
        else:
            print(f"\n✅ スクレイプ完了: {len(jobs)} 件取得（重複チェック前）")
            save_to_excel(jobs)
            # CSV追記（logs/案件検索結果.csv → sync-sheet.mjs で同期）
            csv_path = os.path.join(_ROOT, "logs", "案件検索結果.csv")
            added = save_to_csv(jobs, csv_path)
            print(f"📋 スプシ同期対象: {added} 件追加（{len(jobs) - added} 件は既存CSVと重複）")
            if added > 0:
                print(f"   次のコマンドでスプシに同期してください：")
                print(f"   node server/scripts/sync-sheet.mjs")

    except KeyboardInterrupt:
        print("\n\n⚠️  ユーザーによって中断されました。")

    except Exception as e:
        print("\n" + "=" * 50)
        print("❌ エラーが発生しました")
        print("=" * 50)
        print(f"エラー種別: {type(e).__name__}")
        print(f"エラー内容: {e}")
        print("\n--- 詳細（スタックトレース）---")
        traceback.print_exc()

    finally:
        print("\n" + "=" * 50)
        print("✅ スクリプト実行完了")


if __name__ == "__main__":
    main()
