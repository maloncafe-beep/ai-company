# ランサーズ 案件取得スクリプト（tools/job-scout 版）
#
# 【事前準備】
# pip install playwright openpyxl
# playwright install chromium
#
# 【検索条件の変更方法】
# ランサーズで絞り込み検索した後のURLをコピーして SEARCH_URL に貼り替えるだけでOK
#
# 【ログイン済みChromeで使う場合】
# Chromeをこのコマンドで起動してからランサーズにログイン:
#   "C:/Program Files/Google/Chrome/Application/chrome.exe" --remote-debugging-port=9222
# その後 USE_LOGIN_CHROME = True に変えて実行

import asyncio
import traceback
import csv
import re
from datetime import datetime
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.async_api import async_playwright

# ==================================================
# 設定（ここだけ変更すればOK）
# ==================================================

# ランサーズで絞り込んだ検索結果URLをそのまま貼り替える
SEARCH_URL = (
    "https://www.lancers.jp/work/search"
    "?keyword=%E3%83%A9%E3%82%A4%E3%82%BF%E3%83%BC"  # ライター
    "&open=1"                   # 募集中のみ
    "&sort=new"                 # 新着順
)

# 取得するページ数（1ページ約20件）
MAX_PAGES = 5

# tools/job-scout/ から2階層上がってプロジェクトルート → logs/ へ
_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
LOGS_DIR = os.path.join(_ROOT, "logs", "Case_list_py")
OUTPUT_FILE_XLSX = os.path.join(LOGS_DIR, "lancers_cases_{}.xlsx".format(datetime.now().strftime("%Y%m%d_%H%M%S")))
OUTPUT_FILE_MD = os.path.join(LOGS_DIR, "lancers_cases_latest.md")

# ログイン済みChromeを使う場合は True
USE_LOGIN_CHROME = True
CHROME_DEBUG_PORT = 9222

# ==================================================
# スクレイピング
# ==================================================

async def scrape(url, max_pages):
    results = []

    async with async_playwright() as p:
        try:
            if USE_LOGIN_CHROME:
                print("ログイン済みChromeに接続中...")
                browser = await p.chromium.connect_over_cdp(
                    "http://127.0.0.1:{}".format(CHROME_DEBUG_PORT)
                )
                context = browser.contexts[0]
                page = await context.new_page()
                print("接続OK")
            else:
                browser = await p.chromium.launch(headless=True)
                context = await browser.new_context(
                    user_agent=(
                        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 (KHTML, like Gecko) "
                        "Chrome/124.0.0.0 Safari/537.36"
                    ),
                    viewport={"width": 1280, "height": 800},
                )
                page = await context.new_page()
                print("ブラウザ起動OK")

        except Exception:
            print("【エラー】ブラウザ起動失敗:")
            traceback.print_exc()
            return results

        for page_num in range(1, max_pages + 1):
            paged_url = "{}&page={}".format(url, page_num) if page_num > 1 else url
            print("[{}/{}] 取得中...".format(page_num, max_pages))

            try:
                await page.goto(paged_url, wait_until="domcontentloaded", timeout=60000)
                await page.wait_for_timeout(7000)
            except Exception:
                print("【エラー】ページ読み込み失敗:")
                traceback.print_exc()
                break

            # ポップアップをJSで直接閉じる
            closed = await page.evaluate("""() => {
                const sels = [
                    'button[aria-label="close"]',
                    'button[aria-label="閉じる"]',
                    '.c-modal__close',
                    '[class*="modal"] [class*="close"]',
                    '[class*="dialog"] [class*="close"]',
                ];
                for (const s of sels) {
                    const el = document.querySelector(s);
                    if (el) { el.click(); return s; }
                }
                return null;
            }""")
            if closed:
                print("  -> ポップアップを閉じました")
                await page.wait_for_timeout(600)
            await page.keyboard.press("Escape")
            await page.wait_for_timeout(400)

            # 無限スクロールで全件ロード
            prev_count = 0
            for _ in range(15):
                await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                await page.wait_for_timeout(2000)
                count = await page.eval_on_selector_all(
                    "a",
                    "els => els.filter(e => /lancers\\.jp\\/work\\/detail\\/\\d+/.test(e.href)).length"
                )
                if count == prev_count:
                    break
                prev_count = count

            links = await page.eval_on_selector_all(
                "a",
                """els => [...new Map(
                    els
                    .filter(e => /lancers\\.jp\\/work\\/detail\\/\\d+/.test(e.href))
                    .map(e => [e.href, {href: e.href, text: e.innerText.trim()}])
                ).values()]"""
            )
            print("  -> {} 件のリンクを発見".format(len(links)))

            if not links:
                txt = await page.evaluate("() => document.body.innerText")
                print("  -> ページ内容(先頭200文字): {}".format(txt[:200]))
                break

            seen = set()
            for lnk in links:
                href = lnk["href"]
                if not href or href in seen:
                    continue
                seen.add(href)

                try:
                    info = await page.evaluate("""(href) => {
                        const a = Array.from(document.querySelectorAll('a'))
                                   .find(el => el.href === href);
                        if (!a) return {};
                        const card = a.closest('li,article,[class*="item"],[class*="card"],[class*="work"]')
                                     || a.parentElement;
                        const get = (el, sels) => {
                            for (const s of sels) {
                                try {
                                    const f = el.querySelector(s);
                                    if (f && f.innerText.trim()) return f.innerText.trim();
                                } catch(e) {}
                            }
                            return '';
                        };
                        return {
                            title:     get(card,['h3','h2','h4','[class*="title"]']) || a.innerText.trim(),
                            budget:    get(card,['[class*="budget"]','[class*="price"]','[class*="reward"]']),
                            category:  get(card,['[class*="category"]','[class*="genre"]']),
                            work_type: get(card,['[class*="type"]','[class*="label"]','[class*="badge"]']),
                            posted_at: get(card,['time','[datetime]','[class*="date"]']),
                            client:    get(card,['[class*="client"]','[class*="user"]']),
                        };
                    }""", href)
                except Exception:
                    info = {}

                results.append({
                    "タイトル":     info.get("title", lnk["text"]),
                    "URL":          href,
                    "予算":         info.get("budget", ""),
                    "カテゴリ":     info.get("category", ""),
                    "仕事スタイル": info.get("work_type", ""),
                    "掲載日":       info.get("posted_at", ""),
                    "クライアント": info.get("client", ""),
                    "取得日時":     datetime.now().strftime("%Y-%m-%d %H:%M"),
                })

            print("  -> {} 件取得（累計: {} 件）".format(len(seen), len(results)))
            await asyncio.sleep(2)

        try:
            await browser.close()
        except Exception:
            pass

    return results


# ==================================================
# CSV追記（logs/案件検索結果.csv → sync-sheet.mjs で同期）
# ==================================================

def clean(val):
    """改行・桁区切りカンマ・余分な空白を除去"""
    v = (val or "").replace("\n", " ").replace("\r", "").strip()
    return re.sub(r"(?<=\d),(?=\d)", "", v)


def save_to_csv(jobs, csv_path):
    """
    logs/案件検索結果.csv の形式（11列）に変換して追記する。
    既存IDと重複する案件はスキップ。
    """
    existing_ids = set()
    if os.path.exists(csv_path):
        with open(csv_path, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            next(reader, None)
            for row in reader:
                if row:
                    existing_ids.add(row[0].strip())

    new_rows = []

    for job in jobs:
        url = job.get("URL", "")
        m = re.search(r"/work/(?:detail/)?(\d+)", url)
        job_id = m.group(1) if m else ""
        if not job_id or job_id in existing_ids:
            continue

        budget = clean(job.get("予算", "")) or "要確認"
        # ランサーズは締切が取れないことが多いのでデフォルト日付
        deadline = "2026/09/30"
        category = clean(job.get("カテゴリ", ""))
        work_type = clean(job.get("仕事スタイル", ""))
        memo_parts = [x for x in [category, work_type] if x]
        memo = " / ".join(memo_parts) if memo_parts else ""

        row = [
            job_id,                                        # A: 管理ID
            "ＬＣ",                                         # B: 媒体
            clean(job.get("タイトル", "")),                 # C: 案件タイトル
            budget,                                         # D: 報酬
            clean(job.get("クライアント", "")) or "不明",   # E: クライアント名
            deadline,                                       # F: 納期
            url,                                            # G: URL
            memo,                                           # H: メモ
            "",                                             # I: 検討（社長記入欄）
            "",                                             # J: 仕事の詳細
            "",                                             # K: Masterメモ
        ]
        new_rows.append(row)
        existing_ids.add(job_id)

    if new_rows:
        with open(csv_path, "a", encoding="utf-8-sig", newline="\n") as f:
            writer = csv.writer(f, lineterminator="\n")
            writer.writerows(new_rows)
        print(f"✅ CSV追記完了: {len(new_rows)} 件 → {csv_path}")
    else:
        print("ℹ️  新規案件なし（全件重複またはID未取得）")

    return len(new_rows)


# ==================================================
# Excel出力
# ==================================================

def save_excel(data, path):
    print("\nExcel保存中...")
    wb = Workbook()
    ws = wb.active
    ws.title = "案件一覧"
    headers = ["管理ID", "媒体", "タイトル", "予算", "クライアント", "納期", "URL", "掲載日", "カテゴリ", "仕事スタイル", "取得日時"]
    fill = PatternFill("solid", start_color="13366A", end_color="13366A")
    hfont = Font(bold=True, color="FFFFFF", name="Arial")

    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.fill = fill
        c.font = hfont
        c.alignment = Alignment(horizontal="center", vertical="center")

    if data:
        for ri, row in enumerate(data, 2):
            url = row.get("URL", "")
            m = re.search(r"/work/(?:detail/)?(\d+)", url)
            job_id = m.group(1) if m else ""
            values = [
                job_id,
                "ＬＣ",
                row.get("タイトル", ""),
                row.get("予算", ""),
                row.get("クライアント", ""),
                "2026/09/30",
                url,
                row.get("掲載日", ""),
                row.get("カテゴリ", ""),
                row.get("仕事スタイル", ""),
                row.get("取得日時", ""),
            ]
            for ci, val in enumerate(values, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.alignment = Alignment(vertical="center", wrap_text=(ci == 3))
                if ci == 7 and val:
                    c.hyperlink = val
                    c.font = Font(color="0563C1", underline="single", name="Arial")
    else:
        ws.cell(row=2, column=1, value="案件を取得できませんでした。")

    for ci, w in enumerate([12, 6, 50, 20, 20, 12, 60, 12, 20, 20, 18], 1):
        ws.column_dimensions[ws.cell(1, ci).column_letter].width = w

    ws.row_dimensions[1].height = 25
    ws.freeze_panes = "A2"

    try:
        wb.save(path)
        print("保存完了: {} ({} 件)".format(path, len(data)))
    except Exception:
        print("【エラー】Excel保存失敗:")
        traceback.print_exc()


def save_markdown(data, path):
    print("Markdown保存中...")
    md_lines = [
        "# ランサーズ 案件リスト\n",
        f"**取得日時**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n",
        f"**取得件数**: {len(data)} 件\n",
        "---\n",
        "## 案件一覧\n",
        "| # | タイトル | 予算 | クライアント | URL |\n",
        "|---|------|------|----------|------|\n",
    ]
    for idx, job in enumerate(data, start=1):
        title = job.get("タイトル", "").replace("|", "\\|")[:50]
        budget = job.get("予算", "")
        client = job.get("クライアント", "")[:20]
        url = job.get("URL", "")
        md_lines.append(f"| {idx} | {title} | {budget} | {client} | [Link]({url}) |\n")

    try:
        with open(path, "w", encoding="utf-8") as f:
            f.writelines(md_lines)
        print("Markdown保存完了: {} ({} 件)".format(path, len(data)))
    except Exception:
        print("【エラー】Markdown保存失敗:")
        traceback.print_exc()


# ==================================================
# メイン
# ==================================================

async def main():
    print("=== ランサーズ 案件取得 ===")
    print("モード: {}".format("ログイン済みChrome" if USE_LOGIN_CHROME else "通常（未ログイン）"))
    print("出力先: {}".format(LOGS_DIR))
    print("")

    try:
        data = await scrape(SEARCH_URL, MAX_PAGES)
    except Exception:
        print("【エラー】scrape()で例外発生:")
        traceback.print_exc()
        data = []

    seen = set()
    unique = []
    for item in data:
        if item["URL"] not in seen:
            seen.add(item["URL"])
            unique.append(item)

    print("\n合計 {} 件".format(len(unique)))
    save_excel(unique, OUTPUT_FILE_XLSX)
    save_markdown(unique, OUTPUT_FILE_MD)

    # CSV追記（logs/案件検索結果.csv → sync-sheet.mjs で同期）
    csv_path = os.path.join(_ROOT, "logs", "案件検索結果.csv")
    added = save_to_csv(unique, csv_path)
    print(f"\n📋 スプシ同期対象: {added} 件追加済み")
    print(f"   次のコマンドでスプシに同期してください：")
    print(f"   node server/scripts/sync-sheet.mjs")

    print("\n===完了===")

if __name__ == "__main__":
    asyncio.run(main())
